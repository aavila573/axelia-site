import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()

# Paleta AxelIA
PURPLE = "6C5CE7"
PURPLE2 = "A78BFA"
GREEN = "00D2A0"
DARK = "0A0A0F"
GRAY = "8888A0"
WHITE = "FFFFFF"

header_fill = PatternFill(start_color=PURPLE, end_color=PURPLE, fill_type="solid")
sub_fill = PatternFill(start_color="16161F", end_color="16161F", fill_type="solid")
green_fill = PatternFill(start_color=GREEN, end_color=GREEN, fill_type="solid")
header_font = Font(bold=True, color=WHITE, size=12)
title_font = Font(bold=True, color=WHITE, size=16)
sub_font = Font(bold=True, color=PURPLE2, size=11)
normal_font = Font(color="000000", size=11)
thin = Side(style="thin", color="444455")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

# ============ HOJA 1: Instrucciones ============
ws = wb.active
ws.title = "Instrucciones"
ws.column_dimensions["A"].width = 90
ws["A1"] = "AXELIA — Recolección de datos para el Piloto de 7 días"
ws["A1"].font = title_font
ws["A1"].fill = PatternFill(start_color=DARK, end_color=DARK, fill_type="solid")
ws["A1"].font = Font(bold=True, color=PURPLE2, size=16)

instrucciones = [
    "",
    "CÓMO LLENAR ESTA HOJA (30 minutos máximo)",
    "",
    "1. Hoja 'Datos del negocio': pon tu nombre, WhatsApp y los idiomas en que atiendes.",
    "2. Hoja 'Tours': escribe tus 3 tours MÁS VENDIDOS. No necesitas más para el piloto.",
    "3. Hoja 'Preguntas frecuentes': anota las preguntas que más te hacen a diario.",
    "   (Si no se te ocurren, piensa en las últimas 10 conversaciones de WhatsApp.)",
    "",
    "Con esto, montamos tu agente en 5 días. No tienes que saber nada de tecnología.",
    "",
    "IMPORTANTE:",
    "- Escribe los precios tal cual los cobras hoy (COP).",
    "- Si un tour tiene temporada alta/baja, pon ambos precios.",
    "- En 'Qué incluye' sé específico (ej: 'transporte en lancha, chaleco, guía').",
    "- Si atiendes en varios idiomas, márcalo en 'Idiomas'.",
    "",
    "Cuando termines, guárdalo y envíaselo a tu contacto de AxelIA.",
]
for i, line in enumerate(instrucciones, start=2):
    ws.cell(row=i, column=1, value=line).font = normal_font

# ============ HOJA 2: Datos del negocio ============
ws2 = wb.create_sheet("Datos del negocio")
ws2.column_dimensions["A"].width = 35
ws2.column_dimensions["B"].width = 50
ws2["A1"] = "DATOS DEL NEGOCIO"
ws2["A1"].font = title_font
ws2["A1"].fill = PatternFill(start_color=DARK, end_color=DARK, fill_type="solid")
ws2["A1"].font = Font(bold=True, color=PURPLE2, size=14)

campos = [
    ("Nombre del negocio", ""),
    ("Nombre del contacto", ""),
    ("WhatsApp de atención (número)", ""),
    ("Correo", ""),
    ("Ciudad / Ubicación", ""),
    ("Idiomas en que atiendes hoy", "Español / Inglés / Portugués (marca los que apliquen)"),
    ("Horario de atención actual", "Ej: 8:00 AM - 6:00 PM"),
    ("Temporada alta (fechas)", "Ej: diciembre - enero"),
    ("Temporada baja (fechas)", "Ej: marzo - mayo"),
    ("¿Cómo cobras los anticipos hoy?", "Ej: Nequi, transferencia, en persona"),
]
for i, (campo, ejemplo) in enumerate(campos, start=3):
    ws2.cell(row=i, column=1, value=campo).font = sub_font
    ws2.cell(row=i, column=1).border = border
    c = ws2.cell(row=i, column=2, value=ejemplo)
    c.font = Font(color=GRAY, size=11, italic=True)
    c.border = border

# ============ HOJA 3: Tours ============
ws3 = wb.create_sheet("Tours")
headers = ["Nombre del tour", "Precio (COP)", "Duración", "Horarios de salida", "Cupo máximo", "Qué incluye", "Qué NO incluye", "Notas"]
for i, h in enumerate(headers, start=1):
    ws3.cell(row=1, column=i, value=h)
style_header(ws3, 1, len(headers))
widths = [22, 14, 12, 20, 12, 40, 30, 25]
for i, w in enumerate(widths, start=1):
    ws3.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
# 3 filas de ejemplo vacías para llenar
for r in range(2, 5):
    ws3.cell(row=r, column=1, value="").border = border
    for c in range(2, len(headers) + 1):
        ws3.cell(row=r, column=c).border = border
    ws3.cell(row=r, column=1).fill = sub_fill

# Ejemplos guía
ws3["A6"] = "EJEMPLO (borra y pon los tuyos):"
ws3["A6"].font = Font(bold=True, color=GREEN, size=11)
ejemplo = ["Tour a Johnny Cay", "$90.000", "4 horas", "9:00 AM y 2:00 PM", "30", "Transporte en lancha, guía, almuerzo típico", "Bebidas, snorkel extra", "Salida desde el muelle principal"]
for i, v in enumerate(ejemplo, start=1):
    c = ws3.cell(row=7, column=i, value=v)
    c.font = Font(color=GRAY, size=10, italic=True)
    c.border = border

# ============ HOJA 4: Preguntas frecuentes ============
ws4 = wb.create_sheet("Preguntas frecuentes")
ws4.column_dimensions["A"].width = 45
ws4.column_dimensions["B"].width = 55
ws4["A1"] = "PREGUNTA"
ws4["B1"] = "RESPUESTA"
style_header(ws4, 1, 2)
# 15 filas vacías
for r in range(2, 17):
    ws4.cell(row=r, column=1).border = border
    ws4.cell(row=r, column=2).border = border

ejemplos_faq = [
    ("¿Hay medusas / aguas malas?", "Depende de la temporada. En [mes] suele haber. Recomendamos [recomendación]."),
    ("¿Cuánto dura el tour?", "[X] horas, incluyendo el transporte."),
    ("¿Cómo reservo y pago?", "Reservas por WhatsApp y pagas un anticipo por [Nequi/transferencia]. El resto el día del tour."),
    ("¿Qué pasa si llueve?", "Los tours salen igual salvo condiciones de seguridad. Si cancelamos, te devolvemos el anticipo."),
    ("¿Atienden en inglés/portugués?", "Sí, nuestro equipo habla [idiomas]."),
]
ws4["A19"] = "EJEMPLOS (borra y pon los tuyos):"
ws4["A19"].font = Font(bold=True, color=GREEN, size=11)
for i, (preg, resp) in enumerate(ejemplos_faq, start=20):
    c1 = ws4.cell(row=i, column=1, value=preg)
    c2 = ws4.cell(row=i, column=2, value=resp)
    c1.font = Font(color=GRAY, size=10, italic=True)
    c2.font = Font(color=GRAY, size=10, italic=True)
    c1.border = border
    c2.border = border

wb.save("C:/Users/alexa/AxelIA_Recoleccion_Datos_Piloto.xlsx")
print("✅ Excel creado: AxelIA_Recoleccion_Datos_Piloto.xlsx")
